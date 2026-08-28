"""
Run any time you want a human-readable, color-grouped view of the
current master_flat_table.csv:
    python scripts/export_to_xlsx.py
Produces data/master_flat_table.xlsx.
"""
import csv
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from src.flat_table_manager import CSV_PATH, ALL_COLUMNS

GROUPS = {
    "A_IDENTITY (from M1)": ALL_COLUMNS[0:8],
    "B_CLINICAL_SOURCE (your HIS)": ALL_COLUMNS[8:43],
    "C_CARE_CONTEXT (M2 Phase 1)": ALL_COLUMNS[43:51],
    "D_CONSENT_AND_TRANSFER (M2 Phase 2)": ALL_COLUMNS[51:57],
}
GROUP_COLORS = {
    "A_IDENTITY (from M1)": "2F5496",
    "B_CLINICAL_SOURCE (your HIS)": "833C00",
    "C_CARE_CONTEXT (M2 Phase 1)": "548235",
    "D_CONSENT_AND_TRANSFER (M2 Phase 2)": "7030A0",
}
FONT = "Arial"


def main():
    col_to_group = {c: g for g, cols in GROUPS.items() for c in cols}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Master_Flat_Table"

    col_idx = 1
    for g, cols in GROUPS.items():
        start, end = col_idx, col_idx + len(cols) - 1
        ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)
        c = ws.cell(row=1, column=start, value=g)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill(start_color=GROUP_COLORS[g], end_color=GROUP_COLORS[g], fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center")
        col_idx = end + 1

    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for idx, col in enumerate(ALL_COLUMNS, start=1):
        c = ws.cell(row=2, column=idx, value=col)
        g = col_to_group[col]
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=9)
        c.fill = PatternFill(start_color=GROUP_COLORS[g], end_color=GROUP_COLORS[g], fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
    ws.freeze_panes = "A3"
    ws.row_dimensions[2].height = 40

    if os.path.exists(CSV_PATH):
        with open(CSV_PATH, newline="") as f:
            data_rows = list(csv.DictReader(f))
    else:
        data_rows = []

    for r_idx, row in enumerate(data_rows, start=3):
        for c_idx, col in enumerate(ALL_COLUMNS, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=row.get(col, ""))
            cell.font = Font(name=FONT, size=9)
            cell.border = border

    for idx in range(1, len(ALL_COLUMNS) + 1):
        ws.column_dimensions[get_column_letter(idx)].width = 14

    out_path = os.path.join(os.path.dirname(CSV_PATH), "master_flat_table.xlsx")
    wb.save(out_path)
    print(f"Saved {out_path}")


if __name__ == "__main__":
    main()
